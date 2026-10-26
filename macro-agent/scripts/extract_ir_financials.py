"""손해보험사 IR데이터 엑셀(현대해상 Factsheet 양식: (1) Highlights~(6) BS 시트) → 반기 재무제표 요약 JSON.

python scripts/extract_ir_financials.py --xlsx processed/filings/현대해상_2026Q2_IR데이터.xlsx \
    --company 현대해상 --period 1H26 --avg-assets 46304.9
→ knowledge/portfolios/{회사}/financials_{기간}.json (단위 십억원 = 백만원 ÷ 1,000, value_flag 공시)
행 라벨·열 헤더('26.2Q 누계', '2026. 6월 말' 등)로 값을 찾는다. 다른 회사·양식이면 라벨을 맞춰야 한다.
--avg-assets: 연환산 수익률 분모(평균 운용자산, 십억원) — income_{기간}.json의 avg_invested_assets_krw_bn.
"""
import argparse, json, pathlib
import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("--xlsx", required=True); ap.add_argument("--company", required=True)
ap.add_argument("--period", required=True); ap.add_argument("--avg-assets", type=float, required=True)
args = ap.parse_args()
SRC = args.xlsx
wb = openpyxl.load_workbook(SRC, data_only=True)
norm = lambda s: str(s).replace(' ', '').replace('\n', '').replace('　', '').lstrip('-')
def cols(ws):
    """'26.2Q누계'·'25.2Q누계'·'2026.6월말'·'2025.12월말' 헤더 열 위치"""
    out = {}
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, str):
                v = norm(c.value)
                for k in ("26.2Q누계", "25.2Q누계", "2026.6월말", "2025.12월말", "26.2Q", "25.2Q"):
                    if v == k and k not in out: out[k] = c.column
    return out
def row(ws, label, nth=1, label_cols=(1, 2, 3, 4)):
    n = 0
    for r in ws.iter_rows():
        labs = [norm(c.value) for c in r if c.column in label_cols and isinstance(c.value, str)]
        if norm(label) in labs:
            n += 1
            if n == nth: return {c.column: c.value for c in r}
    raise KeyError(label)
k = lambda v: None if v is None else round(v / 1000, 1)
def flow(sheet, labels):
    ws = wb[sheet]; c = cols(ws); out = {}
    for lab in labels:
        name, nth = (lab if isinstance(lab, tuple) else (lab, 1))
        r = row(ws, name, nth); out[name if nth == 1 else f"{name}#{nth}"] = {"1H26": k(r.get(c["26.2Q누계"])), "1H25": k(r.get(c["25.2Q누계"]))}
    return out
def stock(sheet, labels):
    ws = wb[sheet]; c = cols(ws); out = {}
    for lab in labels:
        r = row(ws, lab); out[lab] = {"2026-06": k(r.get(c["2026.6월말"])), "2025-12": k(r.get(c["2025.12월말"]))}
    return out
fin = {
 "company": args.company, "period": args.period, "unit": "십억원", "value_flag": "공시",
 "source": f"{SRC} (백만원 ÷ 1,000); 외부감사 전 자료",
 "income_statement": flow("(5) IS", ["보험손익", "보험수익", "보험서비스비용", "투자손익", "투자수익", "투자비용", "영업이익", "영업외손익", "세전이익", "법인세비용", "당기순이익"]),
 "uw_general": flow("(3) UW Income", [("보험료배분법적용수익", 1), ("발생손해액", 1), ("사업비", 1), ("보험손익", 1)]),
 "uw_long_term": flow("(3) UW Income", ["보험계약마진상각(+)", "위험조정변동(+)", "예상대비실제차이(+)", "보험금예실차", "예상손해액", ("발생손해액", 2), "사업비예실차", "기타재보험손익(+)", "손실부담계약관련비용(-)", "기타사업비(-)", ("보험손익", 2)]),
 "uw_auto": flow("(3) UW Income", [("보험료배분법적용수익", 2), ("발생손해액", 3), ("사업비", 2), ("보험손익", 3)]),
 "investment": flow("(4) Investment", ["보험금융수익", "이자/배당", "보험금융비용", ("유가증권평가/처분등", 1), ("유가증권평가/처분등", 2), ("파생/외환", 1), ("파생/외환", 2), "투자손익(보험금융손익제외)"]),
 "balance_sheet": stock("(6) BS", ["운용자산", "당기손익-공정가치측정유가증권", "기타포괄손익-공정가치측정유가증권", "상각후원가측정유가증권", "대출채권", "총자산",
                                   "보험계약부채", "최선추정", "위험조정", "보험계약마진", "보험료배분접근법적용", "발생사고요소", "투자계약부채", "기타부채", "부채총계",
                                   "이익잉여금", "해약환급금준비금", "기타포괄손익", "자본총계"]),
}
fin["sales"] = flow("(1) Highlights", ["일반", "장기(신계약)", "인보험(월납환산)", "자동차"])
fin["sales"]["note"] = "장기(신계약)·인보험은 월납환산 신계약, 일반·자동차는 원수보험료"
# 파생 지표
li, ls = fin["uw_long_term"], fin["uw_general"]; au = fin["uw_auto"]
def ratio(a, b): return round(a / b * 100, 1)
fin["derived"] = {
 "일반_손해율_pct": {p: ratio(fin["uw_general"]["발생손해액"][p], fin["uw_general"]["보험료배분법적용수익"][p]) for p in ("1H26", "1H25")},
 "일반_사업비율_pct": {p: ratio(fin["uw_general"]["사업비"][p], fin["uw_general"]["보험료배분법적용수익"][p]) for p in ("1H26", "1H25")},
 "자동차_손해율_pct": {p: ratio(au["발생손해액#3"][p], au["보험료배분법적용수익#2"][p]) for p in ("1H26", "1H25")},
 "자동차_사업비율_pct": {p: ratio(au["사업비#2"][p], au["보험료배분법적용수익#2"][p]) for p in ("1H26", "1H25")},
 "장기_발생/예상손해_pct": {p: ratio(li["발생손해액#2"][p], li["예상손해액"][p]) for p in ("1H26", "1H25")},
 "보험금융손익": {p: round(fin["investment"]["보험금융수익"][p] - fin["investment"]["보험금융비용"][p], 1) for p in ("1H26", "1H25")},
}
fin["derived"]["자동차_합산비율_pct"] = {p: round(fin["derived"]["자동차_손해율_pct"][p] + fin["derived"]["자동차_사업비율_pct"][p], 1) for p in ("1H26", "1H25")}
avg = args.avg_assets
fin["derived"]["연환산_운용자산대비_pct_1H26"] = {
 "투자손익(보험금융손익제외)": round(fin["investment"]["투자손익(보험금융손익제외)"]["1H26"] * 2 / avg * 100, 2),
 "보험금융비용": round(fin["investment"]["보험금융비용"]["1H26"] * 2 / avg * 100, 2),
 "보험금융손익(순)": round(fin["derived"]["보험금융손익"]["1H26"] * 2 / avg * 100, 2),
 "투자손익": round(fin["income_statement"]["투자손익"]["1H26"] * 2 / avg * 100, 2),
 "기준": f"평균 운용자산 {avg:,.1f}십억원"}
# 장기 보험손익 증감 분해
fin["derived"]["장기보험손익_증감분해"] = {lab: round(li[lab]["1H26"] - li[lab]["1H25"], 1) * (-1 if lab in ("손실부담계약관련비용(-)", "기타사업비(-)") else 1)
                                    for lab in ["보험계약마진상각(+)", "위험조정변동(+)", "예상대비실제차이(+)", "기타재보험손익(+)", "손실부담계약관련비용(-)", "기타사업비(-)"]}
fin["derived"]["장기보험손익_증감합계"] = round(li["보험손익#2"]["1H26"] - li["보험손익#2"]["1H25"], 1)
p = pathlib.Path(f"knowledge/portfolios/{args.company}/financials_{args.period}.json")
p.write_text(json.dumps(fin, ensure_ascii=False, indent=2) + "\n")
print(f"저장: {p}")
